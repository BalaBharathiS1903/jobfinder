import uuid

from rest_framework import generics, permissions
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework_simplejwt.exceptions import TokenError
from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.mail import send_mail
from django.core.validators import validate_email
from django.db import IntegrityError
from django.contrib.auth.password_validation import validate_password
from .serializers import RegisterSerializer, UserSerializer


class RegisterView(generics.CreateAPIView):
    serializer_class = RegisterSerializer
    permission_classes = (permissions.AllowAny,)


class MeView(generics.RetrieveAPIView):
    serializer_class = UserSerializer

    def get_object(self):
        return self.request.user


def _set_auth_cookies(response, access, refresh):
    secure = not settings.DEBUG
    response.set_cookie("access", access, httponly=True, secure=secure, samesite="Lax", max_age=3600)
    response.set_cookie("refresh", refresh, httponly=True, secure=secure, samesite="Lax", max_age=604800)


@api_view(["POST"])
@permission_classes([AllowAny])
def login_view(request):
    from django.contrib.auth import authenticate
    email = request.data.get("email", "").strip()
    password = request.data.get("password", "")
    user = authenticate(request, email=email, password=password)
    if not user:
        return Response({"error": "Invalid credentials."}, status=401)
    refresh = RefreshToken.for_user(user)
    resp = Response({
        "detail": "Login successful.",
        "access": str(refresh.access_token),
        "refresh": str(refresh),
    })
    _set_auth_cookies(resp, str(refresh.access_token), str(refresh))
    return resp


@api_view(["POST"])
@permission_classes([AllowAny])
def refresh_view(request):
    token = request.COOKIES.get("refresh") or request.data.get("refresh", "")
    if not token:
        return Response({"error": "No refresh token."}, status=401)
    try:
        refresh = RefreshToken(token)
        access = str(refresh.access_token)
        resp = Response({"detail": "Refreshed.", "access": access})
        _set_auth_cookies(resp, access, str(refresh))
        return resp
    except TokenError:
        return Response({"error": "Invalid or expired refresh token."}, status=401)


@api_view(["POST"])
@permission_classes([AllowAny])
def logout_view(request):
    resp = Response({"detail": "Logged out."})
    resp.delete_cookie("access")
    resp.delete_cookie("refresh")
    return resp


@api_view(["POST"])
@permission_classes([AllowAny])
def forgot_password_view(request):
    email = request.data.get("email", "").strip().lower()
    if not email:
        return Response({"error": "Email is required."}, status=400)
    try:
        validate_email(email)
    except ValidationError:
        return Response({"error": "Enter a valid email address."}, status=400)
    from apps.accounts.models import PasswordResetToken, User

    detail = "If that email exists, a password reset link has been sent."
    response_data = {"detail": detail}

    try:
        user = User.objects.get(email__iexact=email)
    except User.DoesNotExist:
        return Response(response_data)

    PasswordResetToken.objects.filter(user=user, used=False).update(used=True)
    reset_token = PasswordResetToken.objects.create(user=user)
    reset_url = f"{settings.FRONTEND_URL.rstrip('/')}/reset-password?token={reset_token.token}"

    send_mail(
        "Reset your JobFinder password",
        (
            "Use the link below to reset your password. "
            "This link expires in 1 hour and can be used once.\n\n"
            f"{reset_url}"
        ),
        settings.DEFAULT_FROM_EMAIL,
        [user.email],
        fail_silently=True,
    )

    if settings.DEBUG:
        response_data.update({"reset_token": str(reset_token.token), "reset_url": reset_url})
    return Response(response_data)


@api_view(["POST"])
@permission_classes([AllowAny])
def validate_default_reset_email_view(request):
    email = request.data.get("email", "").strip().lower()
    if not email:
        return Response({"error": "Email is required."}, status=400)
    try:
        validate_email(email)
    except ValidationError:
        return Response({"error": "Enter a valid email address."}, status=400)

    from apps.accounts.models import User
    if not User.objects.filter(email__iexact=email).exists():
        return Response({"error": "No account found with this email address."}, status=404)
    return Response({"detail": "Email verified."})


@api_view(["POST"])
@permission_classes([AllowAny])
def reset_password_view(request):
    from apps.accounts.models import PasswordResetToken
    token_value = request.data.get("token", "")
    new_password = request.data.get("password", "")
    if not token_value or not new_password:
        return Response({"error": "Reset token and password are required."}, status=400)
    try:
        parsed_token = uuid.UUID(str(token_value))
    except (TypeError, ValueError):
        return Response({"error": "Invalid or expired reset token."}, status=400)
    try:
        reset_token = PasswordResetToken.objects.select_related("user").get(token=parsed_token)
    except PasswordResetToken.DoesNotExist:
        return Response({"error": "Invalid or expired reset token."}, status=400)
    if not reset_token.is_valid():
        return Response({"error": "Invalid or expired reset token."}, status=400)
    try:
        validate_password(new_password, user=reset_token.user)
    except ValidationError as exc:
        return Response({"error": list(exc.messages)}, status=400)

    reset_token.user.set_password(new_password)
    reset_token.user.save(update_fields=["password"])
    reset_token.used = True
    reset_token.save(update_fields=["used"])
    return Response({"detail": "Password has been updated."})


@api_view(["POST"])
@permission_classes([AllowAny])
def default_password_reset_view(request):
    from django.contrib.auth import authenticate
    email = request.data.get("email", "").strip().lower()
    default_password = request.data.get("default_password", "")
    new_password = request.data.get("new_password", "")

    if not email or not default_password or not new_password:
        return Response({"error": "Email, default password, and new password are required."}, status=400)
    try:
        validate_email(email)
    except ValidationError:
        return Response({"error": "Enter a valid email address."}, status=400)
    if default_password != "vdart@#12345":
        return Response({"error": "Default password is incorrect."}, status=400)

    user = authenticate(request, email=email, password=default_password)
    if not user:
        return Response({"error": "This account cannot be reset with the default password."}, status=400)
    try:
        validate_password(new_password, user=user)
    except ValidationError as exc:
        return Response({"error": list(exc.messages)}, status=400)

    user.set_password(new_password)
    user.save(update_fields=["password"])
    return Response({"detail": "Password has been updated. You can now sign in with your new password."})


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def change_password_view(request):
    current_password = request.data.get("current_password", "")
    new_password = request.data.get("new_password", "")
    if not current_password or not new_password:
        return Response({"error": "Current and new password are required."}, status=400)
    user = request.user
    if not user.check_password(current_password):
        return Response({"error": "Current password is incorrect."}, status=400)
    try:
        validate_password(new_password, user=user)
    except ValidationError as exc:
        return Response({"error": list(exc.messages)}, status=400)
    user.set_password(new_password)
    user.save()
    return Response({"detail": "Password changed successfully."})


# ── User activity endpoint ───────────────────────────────────────────────────

@api_view(["GET"])
@permission_classes([IsAuthenticated])
def my_activity(request):
    from apps.resume.models import Resume, ResumeVersion
    from apps.resume.serializers import ResumeSerializer
    from apps.jobs.models import JobSearch, SavedJob
    user = request.user

    resumes = Resume.objects.filter(user=user).order_by("-uploaded_at")
    resume_data = ResumeSerializer(resumes, many=True).data

    # Per-resume reparse/replace history count
    for r in resume_data:
        r["version_count"] = ResumeVersion.objects.filter(resume_id=r["id"]).count()

    job_searches = list(
        JobSearch.objects.filter(user=user).order_by("-searched_at")
        .values("id", "query", "location", "searched_at")
    )

    saved_jobs = list(
        SavedJob.objects.filter(user=user).order_by("-saved_at")
        .values("id", "title", "company", "location", "url", "saved_at")
    )

    # Daily job search count (today)
    from django.utils import timezone
    today_start = timezone.now().replace(hour=0, minute=0, second=0, microsecond=0)
    daily_search_count = JobSearch.objects.filter(user=user, searched_at__gte=today_start).count()

    # Course progress
    from apps.courses.models import CourseProgress, CourseCertificate
    course_progress_qs = CourseProgress.objects.filter(user=user)
    course_progress = []
    for cp in course_progress_qs:
        completed_count = sum(1 for v in cp.completed.values() if v)
        course_progress.append({
            "course_id": cp.course_id,
            "completed_lessons": completed_count,
            "updated_at": cp.updated_at,
        })
    certificates = list(
        CourseCertificate.objects.filter(user=user)
        .values("course_id", "cert_id", "completed_on", "issued_at")
    )

    return Response({
        "limits": {
            "resume_upload_limit": user.resume_upload_limit,
            "job_search_limit": user.job_search_limit,
        },
        "resume_count": resumes.count(),
        "resumes": resume_data,
        "job_search_total": len(job_searches),
        "job_search_today": daily_search_count,
        "job_searches": job_searches,
        "saved_jobs_count": len(saved_jobs),
        "saved_jobs": saved_jobs,
        "course_progress": course_progress,
        "certificates": certificates,
    })


# ── Admin endpoints ───────────────────────────────────────────────────────────

@api_view(["GET"])
@permission_classes([IsAuthenticated])
def admin_users_list(request):
    if not request.user.is_superuser:
        return Response({"error": "Forbidden."}, status=403)
    from django.contrib.auth import get_user_model
    users = get_user_model().objects.all().order_by("-date_joined")
    return Response(UserSerializer(users, many=True).data)


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def admin_create_user(request):
    if not request.user.is_superuser:
        return Response({"error": "Forbidden."}, status=403)
    serializer = RegisterSerializer(data=request.data)
    if serializer.is_valid():
        user = serializer.save()
        return Response(UserSerializer(user).data, status=201)
    return Response(serializer.errors, status=400)


def _truthy(value):
    return str(value).strip().lower() in {"1", "true", "yes", "y", "active", "granted"}


def _coerce_bool(value):
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "y", "active", "granted"}:
        return True
    if normalized in {"0", "false", "no", "n", "inactive", "disabled", "denied"}:
        return False
    raise ValueError("Enter a valid boolean value.")


def _coerce_limit(value):
    try:
        return max(0, int(value))
    except (TypeError, ValueError):
        raise ValueError("Enter a valid non-negative number.")


def _normalize_bulk_key(value):
    return str(value or "").strip().lower().replace(" ", "_")


def _normalize_bulk_row(row):
    return {_normalize_bulk_key(k): v for k, v in row.items()}


def _read_bulk_user_rows(uploaded_file):
    name = uploaded_file.name.lower()
    if name.endswith(".csv"):
        import csv
        import io
        try:
            text = uploaded_file.read().decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("CSV files must be UTF-8 encoded.")
        reader = csv.DictReader(io.StringIO(text))
        headers = [_normalize_bulk_key(h) for h in (reader.fieldnames or [])]
        if not headers:
            return []
        if "email" not in headers:
            raise ValueError("The uploaded file must include an email column.")
        return [_normalize_bulk_row(row) for row in reader]

    if name.endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError("Excel support is not installed. Run pip install -r requirements.txt.")
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [_normalize_bulk_key(h) for h in rows[0]]
        if "email" not in headers:
            raise ValueError("The uploaded file must include an email column.")
        return [
            {headers[i]: value for i, value in enumerate(row) if i < len(headers)}
            for row in rows[1:]
            if any(value not in (None, "") for value in row)
        ]

    raise ValueError("Upload a .xlsx or .csv file.")


@api_view(["POST"])
@permission_classes([IsAuthenticated])
def admin_bulk_create_users(request):
    if not request.user.is_superuser:
        return Response({"error": "Forbidden."}, status=403)

    uploaded_file = request.FILES.get("file")
    if not uploaded_file:
        return Response({"error": "No file uploaded."}, status=400)

    try:
        rows = _read_bulk_user_rows(uploaded_file)
    except ValueError as exc:
        return Response({"error": str(exc)}, status=400)

    from django.contrib.auth import get_user_model
    from django.utils.crypto import get_random_string

    User = get_user_model()
    created = []
    skipped = []

    for index, raw in enumerate(rows, start=2):
        row = _normalize_bulk_row(raw)
        email = str(row.get("email") or "").strip().lower()
        username = str(row.get("username") or row.get("name") or "").strip()
        password = str(row.get("password") or "").strip()

        if not email:
            skipped.append({"row": index, "reason": "Missing email."})
            continue
        try:
            validate_email(email)
        except ValidationError:
            skipped.append({"row": index, "email": email, "reason": "Invalid email."})
            continue
        if User.objects.filter(email__iexact=email).exists():
            skipped.append({"row": index, "email": email, "reason": "Email already exists."})
            continue

        if not username:
            username = email.split("@")[0]
        base_username = username[:140]
        username = base_username
        suffix = 1
        while User.objects.filter(username__iexact=username).exists():
            username = f"{base_username[:135]}{suffix}"
            suffix += 1

        generated_password = False
        if not password:
            password = get_random_string(12)
            generated_password = True

        try:
            user = User.objects.create_user(username=username, email=email, password=password)
        except (IntegrityError, ValueError) as exc:
            skipped.append({"row": index, "email": email, "reason": str(exc)})
            continue

        for field in ("has_prep_access", "is_active"):
            if field in row and row[field] not in (None, ""):
                setattr(user, field, _truthy(row[field]))
        for field in ("resume_upload_limit", "job_search_limit"):
            if field in row and row[field] not in (None, ""):
                try:
                    setattr(user, field, max(0, int(row[field])))
                except (TypeError, ValueError):
                    pass
        user.save()

        item = UserSerializer(user).data
        item["password"] = password if generated_password else ""
        created.append(item)

    return Response({
        "created_count": len(created),
        "skipped_count": len(skipped),
        "created": created,
        "skipped": skipped,
    }, status=201)


@api_view(["GET"])
@permission_classes([IsAuthenticated])
def admin_user_detail(request, pk):
    """Full user data: account + profile + resumes + activity."""
    if not request.user.is_superuser:
        return Response({"error": "Forbidden."}, status=403)
    from django.contrib.auth import get_user_model
    from apps.profile.models import UserProfile
    from apps.profile.serializers import UserProfileSerializer
    from apps.resume.models import Resume
    from apps.resume.serializers import ResumeSerializer
    from apps.jobs.models import JobSearch, SavedJob
    try:
        target = get_user_model().objects.get(pk=pk)
    except get_user_model().DoesNotExist:
        return Response(status=404)
    profile = None
    try:
        profile = UserProfileSerializer(target.profile, context={'request': request}).data
    except UserProfile.DoesNotExist:
        pass
    resumes = ResumeSerializer(Resume.objects.filter(user=target).order_by("-uploaded_at"), many=True).data
    job_searches = list(JobSearch.objects.filter(user=target).order_by("-searched_at").values("query", "location", "searched_at")[:20])
    saved_jobs = list(SavedJob.objects.filter(user=target).order_by("-saved_at").values("title", "company", "location", "saved_at")[:20])
    return Response({
        "user": UserSerializer(target).data,
        "profile": profile,
        "resumes": resumes,
        "job_searches": job_searches,
        "saved_jobs": saved_jobs,
    })


@api_view(["PATCH"])
@permission_classes([IsAuthenticated])
def admin_user_update(request, pk):
    if not request.user.is_superuser:
        return Response({"error": "Forbidden."}, status=403)
    from django.contrib.auth import get_user_model
    try:
        target = get_user_model().objects.get(pk=pk)
    except get_user_model().DoesNotExist:
        return Response(status=404)
    if target.is_superuser and target != request.user:
        return Response({"error": "Cannot modify other superadmin accounts."}, status=400)
    if target == request.user and "is_active" in request.data:
        return Response({"error": "Cannot deactivate yourself."}, status=400)

    if "username" in request.data:
        username = request.data.get("username", "").strip()
        if not username:
            return Response({"username": "Username cannot be blank."}, status=400)
        if username.lower() != target.username.lower() and get_user_model().objects.filter(username__iexact=username).exclude(pk=target.pk).exists():
            return Response({"username": "This username is already taken."}, status=400)
        target.username = username

    if "email" in request.data:
        email = request.data.get("email", "").strip().lower()
        if not email:
            return Response({"email": "Email cannot be blank."}, status=400)
        try:
            validate_email(email)
        except ValidationError:
            return Response({"email": "Enter a valid email address."}, status=400)
        if email != target.email.lower() and get_user_model().objects.filter(email__iexact=email).exclude(pk=target.pk).exists():
            return Response({"email": "An account with this email already exists."}, status=400)
        target.email = email

    for field in ("is_active", "has_prep_access"):
        if field in request.data:
            try:
                setattr(target, field, _coerce_bool(request.data[field]))
            except ValueError as exc:
                return Response({field: str(exc)}, status=400)
    for field in ("resume_upload_limit", "job_search_limit"):
        if field in request.data:
            try:
                setattr(target, field, _coerce_limit(request.data[field]))
            except ValueError as exc:
                return Response({field: str(exc)}, status=400)
    target.save()
    return Response(UserSerializer(target).data)


@api_view(["DELETE"])
@permission_classes([IsAuthenticated])
def admin_user_delete(request, pk):
    if not request.user.is_superuser:
        return Response({"error": "Forbidden."}, status=403)
    from django.contrib.auth import get_user_model
    try:
        target = get_user_model().objects.get(pk=pk)
    except get_user_model().DoesNotExist:
        return Response(status=404)
    if target.is_superuser:
        return Response({"error": "Cannot delete superadmin accounts."}, status=400)
    target.delete()
    return Response(status=204)
